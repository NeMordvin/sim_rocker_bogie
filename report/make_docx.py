"""Собирает Отчёт.docx из Отчёт.md (TNR 12, по ширине, интервал 1,25)."""
import re, pathlib
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A, WD_COLOR_INDEX
from docx.oxml.ns import qn

root = pathlib.Path(__file__).parent
md = (root / 'Отчёт.md').read_text(encoding='utf-8').splitlines()
F = 'Times New Roman'

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)

def setfont(style, size, bold=None):
    style.font.name = F; style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    style.element.rPr.rFonts.set(qn('w:eastAsia'), F)
    style.font.color.rgb = None
    pf = style.paragraph_format
    pf.line_spacing = 1.25; pf.space_before = pf.space_after = Pt(0)

for n in ('Normal', 'List Bullet', 'List Number', 'Caption'):
    setfont(doc.styles[n], 12)
setfont(doc.styles['Title'], 18, True)
setfont(doc.styles['Heading 1'], 14, True)
for n in ('Title', 'Heading 1'):
    st = doc.styles[n]
    st.font.italic = False
    pPr = st.element.get_or_add_pPr()
    for b in pPr.findall(qn('w:pBdr')): pPr.remove(b)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)

TOK = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
def add_runs(p, text, **kw):
    for part in TOK.split(text):
        if not part: continue
        b = i = False
        if part.startswith('**'): part, b = part[2:-2], True
        elif part.startswith('*'): part, i = part[1:-1], True
        elif part.startswith('`'): part = part[1:-1]
        r = p.add_run(part); r.bold = b or None; r.italic = i or None
        for k, v in kw.items(): setattr(r, k, v)

def para(text, style=None, align=A.JUSTIFY, indent=True, **kw):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    if indent: p.paragraph_format.first_line_indent = Cm(1.25)
    add_runs(p, text, **kw)
    return p

num = 0
for l in md:
    if not l.strip(): continue
    if l.startswith('# '):
        p = doc.add_paragraph(l[2:], style='Title'); p.alignment = A.CENTER
    elif l.startswith('## '):
        p = doc.add_paragraph(l[3:], style='Heading 1'); p.alignment = A.LEFT
    elif l.startswith('$$'):
        p = doc.add_paragraph(); p.alignment = A.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
        def t(x, sub=False, it=True):
            r = p.add_run(x); r.italic = it; r.font.subscript = sub
        t('J = Σ w'); t('i', True, False); t(' h'); t('i', True, False)
        t(' → max,   w'); t('ступень', True, False); t(' = 1,  w'); t('съезд', True, False)
        t(' = 0,3,  w'); t('ров', True, False); t(' = w'); t('кочки', True, False); t(' = 0,5.')
    elif l.startswith('>'):
        p = para(re.sub(r'\*\*', '', l[1:].strip()), align=A.CENTER, indent=False)
        for r in p.runs: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
    elif l.startswith('!['):
        p = doc.add_paragraph(); p.alignment = A.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(root / re.search(r'\((.*?)\)', l).group(1)), width=Cm(14.5))
    elif l.startswith('*Рис'):
        p = para(l.strip('*'), align=A.CENTER, indent=False)
        p.paragraph_format.space_after = Pt(8)
    elif l.startswith('- '):
        para(l[2:], style='List Bullet', indent=False)
    elif re.match(r'\d+\.\s', l):
        num += 1
        p = para(f'{num}. ' + re.sub(r'^\d+\.\s', '', l), indent=True)
    else:
        para(l)

out = root / 'Отчёт.docx'
doc.save(out); print(out)
